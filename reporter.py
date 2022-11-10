import openpyxl as ex
from openpyxl.chart import (LineChart, Reference)
import datetime as dt


class ExcelReporter:
    report_book = None
    DEFAULT_SHEET_NAME = "Отчёт"
    __active_sheet = None

    def create_book(self, file_name):
        self.report_book = ex.Workbook()
        self.report_book.create_sheet(self.DEFAULT_SHEET_NAME, index=0)
        self.report_book.remove(self.report_book["Sheet"])
        self.__active_sheet = self.report_book.active
        self.__write_header(file_name)

    def insert_value(self, cell_name, value, sheet_name=None):
        if sheet_name is None:
            sn = self.DEFAULT_SHEET_NAME
        else:
            sn = sheet_name
        sheet = self.report_book[sn]
        cell = sheet[cell_name]
        cell.value = value

    def save_book(self, prg_name, base_path=""):
        now = dt.datetime.now()
        curr_data = now.strftime("%d-%m-%Y %H.%M.%S")
        self.report_book.save(f"{base_path}Отчёт {prg_name} {curr_data}.xlsx")

    def place_linechart(self, chart_cell, min_col, max_col, min_row, max_row, chart_name, x_axis_name, y_axis_name):
        data = Reference(worksheet=self.__active_sheet, min_col=min_col, max_col=max_col, min_row=min_row, max_row=max_row)
        chart = LineChart()
        chart.add_data(data, titles_from_data=True)
        chart.title = chart_name
        chart.x_axis = x_axis_name
        chart.y_axis = y_axis_name
        self.__active_sheet.add_chart(chart, chart_cell)


    def __write_header(self, file_name):
        self.insert_value("A1", value=f"Отчет по файлу исходного кода")
        self.insert_value("B1", value=f"{file_name}")
        self.insert_value("A2", value="Тип")
        self.insert_value("B2", value="Содержание")
        self.insert_value("C2", value="Статус")
        self.insert_value("D2", value="Важность")
        self.insert_value("F1", value="Динамика")
        self.insert_value("F2", value="Неиспользуемые переменные")
        self.insert_value("G2", value="Неправильные имена переменных")
        self.insert_value("H2", value="Неправильные имена директив")
        self.insert_value("I2", value="Непарность скобок")
        self.insert_value("J2", value="Время проверки")
